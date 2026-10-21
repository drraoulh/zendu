#!/usr/bin/env python3
"""
Cahier de tests Sandbox Orange Money (guide d'intégration v1.0, section 14.1, scénarios T01 à T10).

Les identifiants viennent UNIQUEMENT des variables d'environnement (jamais du code) :
  OM_CONSUMER_KEY, OM_CONSUMER_SECRET   clés Sandbox de l'application (plateforme OM-API Manager)
  OM_X_AUTH_TOKEN, OM_CHANNEL_MSISDN, OM_PIN   paramètres Sandbox transmis par Orange Money
  OM_TEST_SUBSCRIBER   numéro Orange Money de test (9 chiffres), par défaut 690000000
  OM_NOTIF_URL         URL de notification (par défaut celle de WorldSoft Transfer)
  OM_FAIL_WAIT         secondes d'attente pour le scénario T06 (par défaut 180)

Résultat : out/cahier-tests-orange-money-<date>.xlsx et .md (aucun secret ni jeton n'y est écrit).
Usage : python3 orange_sandbox_test.py
"""
import os
import sys
import time
import uuid
from datetime import datetime
from pathlib import Path

import requests

HERE = Path(__file__).resolve().parent
TOKEN_URL = os.environ.get("OM_TOKEN_URL", "https://apistore.orange.cm/oauth2/token")
BASE_URL = os.environ.get("OM_BASE_URL", "https://apistore-gateway.orange.cm/omcoreapis/1.0.2").rstrip("/")
NOTIF_URL = os.environ.get("OM_NOTIF_URL", "https://worldsoft-transfer.vercel.app/api/webhooks/orange")
SUBSCRIBER = os.environ.get("OM_TEST_SUBSCRIBER", "690000000")
FAIL_WAIT = int(os.environ.get("OM_FAIL_WAIT", "180"))


def env(name):
    value = os.environ.get(name, "").strip()
    if not value:
        sys.exit(f"Variable d'environnement manquante : {name}")
    return value


CONSUMER_KEY = env("OM_CONSUMER_KEY")
CONSUMER_SECRET = env("OM_CONSUMER_SECRET")
X_AUTH_TOKEN = env("OM_X_AUTH_TOKEN")
CHANNEL = env("OM_CHANNEL_MSISDN")
PIN = env("OM_PIN")

# Le serveur du jeton n'envoie pas son certificat intermédiaire DigiCert : on le fournit (vérification TLS complète).
CA_FILE = HERE / "digicert-global-g2-tls-rsa-sha256-2020-ca1.crt"
if os.environ.get("REQUESTS_CA_BUNDLE"):
    bundle = HERE / "out" / ".ca-bundle.pem"
    bundle.parent.mkdir(exist_ok=True)
    bundle.write_text(Path(os.environ["REQUESTS_CA_BUNDLE"]).read_text() + "\n" + CA_FILE.read_text())
    VERIFY = str(bundle)
else:
    import certifi

    bundle = HERE / "out" / ".ca-bundle.pem"
    bundle.parent.mkdir(exist_ok=True)
    bundle.write_text(Path(certifi.where()).read_text() + "\n" + CA_FILE.read_text())
    VERIFY = str(bundle)

results = []


def now():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def record(num, scenario, endpoint, expected, resp, ok, note="", data=None):
    body = {}
    if resp is not None:
        try:
            body = resp.json()
        except ValueError:
            body = {"raw": resp.text[:200]}
    d = data if data is not None else (body.get("data") or {}) if isinstance(body, dict) else {}
    results.append(
        {
            "N°": num,
            "Scénario": scenario,
            "Endpoint": endpoint,
            "Résultat attendu": expected,
            "HTTP": resp.status_code if resp is not None else "",
            "Message": (body.get("message") if isinstance(body, dict) else "") or note,
            "Statut": d.get("status", "") if isinstance(d, dict) else "",
            "payToken": d.get("payToken", "") if isinstance(d, dict) else "",
            "txnid": d.get("txnid", "") if isinstance(d, dict) else "",
            "orderId": d.get("orderId", "") if isinstance(d, dict) else "",
            "Résultat": "OK" if ok else "KO",
            "Commentaire": note,
            "Date et heure": now(),
        }
    )
    print(f"{num} {'OK' if ok else 'KO'}  HTTP {resp.status_code if resp is not None else '-'}  {note}")


def get_token():
    return requests.post(
        TOKEN_URL,
        auth=(CONSUMER_KEY, CONSUMER_SECRET),
        data={"grant_type": "client_credentials"},
        headers={"Content-Type": "application/x-www-form-urlencoded"},
        timeout=30,
        verify=VERIFY,
    )


def headers(token):
    return {"wso2-auth": f"Bearer {token}", "X-AUTH-TOKEN": X_AUTH_TOKEN, "Content-Type": "application/json"}


def post(path, token, payload=None):
    return requests.post(f"{BASE_URL}{path}", headers=headers(token), json=payload, timeout=60, verify=VERIFY)


def get(path, token):
    return requests.get(f"{BASE_URL}{path}", headers=headers(token), timeout=30, verify=VERIFY)


def pay_body(pay_token, order_id, amount, description):
    return {
        "subscriberMsisdn": SUBSCRIBER,
        "channelUserMsisdn": CHANNEL,
        "amount": str(amount),
        "description": description,
        "orderId": order_id,
        "pin": PIN,
        "payToken": pay_token,
        "notifUrl": NOTIF_URL,
    }


def order_id(prefix):
    return f"WST-{prefix}-{uuid.uuid4().hex[:8].upper()}"


# T01 — jeton d'accès
r = get_token()
token = r.json().get("access_token") if r.status_code == 200 else None
record("T01", "Obtenir un access token", "POST /oauth2/token", "200, access_token et expires_in", r,
       bool(token and r.json().get("expires_in")), f"expires_in = {r.json().get('expires_in')}" if token else "", data={})
if not token:
    print("Jeton refusé : vérifiez OM_CONSUMER_KEY / OM_CONSUMER_SECRET (clés Sandbox).")

if token:
    # T02 — initialisation paiement marchand
    r = post("/mp/init", token)
    pt = (r.json().get("data") or {}).get("payToken") if r.status_code == 200 else None
    record("T02", "Initialiser un paiement marchand", "POST /mp/init", "200, payToken renvoyé", r, bool(pt))

    # T03 — paiement (PENDING attendu en Sandbox)
    oid = order_id("MP")
    r = post("/mp/pay", token, pay_body(pt, oid, 100, "WorldSoft Transfer test T03")) if pt else None
    st = (r.json().get("data") or {}).get("status") if r is not None and r.status_code == 200 else None
    record("T03", "Paiement confirmé par le client", "POST /mp/pay", "PENDING en Sandbox", r, st == "PENDING")

    # T04 — relance de la confirmation
    r = get(f"/mp/push/{pt}", token) if pt else None
    record("T04", "Relance de la confirmation", "GET /mp/push/{payToken}", "200, « Push sent to customer »", r,
           r is not None and r.status_code == 200)

    # T05 — statut
    r = get(f"/mp/paymentstatus/{pt}", token) if pt else None
    record("T05", "Consultation du statut", "GET /mp/paymentstatus/{payToken}", "Statut à jour", r,
           r is not None and r.status_code == 200 and bool((r.json().get("data") or {}).get("status")))

    # T07 — réutilisation du même payToken : doit être refusée (aucun second débit)
    r = post("/mp/pay", token, pay_body(pt, order_id("MP"), 100, "WorldSoft Transfer test T07")) if pt else None
    refused = r is not None and (r.status_code != 200 or (r.json().get("data") or {}).get("status") in ("FAILED", None))
    record("T07", "Réutilisation d'un payToken déjà utilisé", "POST /mp/pay", "Erreur, aucun second débit", r, refused)

    # T08 — jeton expiré/invalide : 401, puis nouveau jeton et succès
    r_bad = post("/mp/init", "jeton-invalide-test-T08")
    r_new = get_token()
    t2 = r_new.json().get("access_token") if r_new.status_code == 200 else None
    r_ok = post("/mp/init", t2) if t2 else None
    record("T08", "Access token expiré", "POST /mp/init", "401, puis nouveau token et succès", r_ok,
           r_bad.status_code == 401 and r_ok is not None and r_ok.status_code == 200,
           f"1er appel avec jeton invalide : HTTP {r_bad.status_code}")
    token = t2 or token

    # T09 — Cashin complet (versement vers un portefeuille Orange Money : usage de WorldSoft Transfer)
    r = post("/cashin/init", token)
    ci = (r.json().get("data") or {}).get("payToken") if r.status_code == 200 else None
    record("T09a", "Cashin : initialisation", "POST /cashin/init", "200, payToken renvoyé", r, bool(ci))
    if ci:
        body = pay_body(ci, order_id("CI"), 500, "WorldSoft Transfer test T09")
        body["amount"] = 500
        r = post("/cashin/pay", token, body)
        st = (r.json().get("data") or {}).get("status") if r.status_code == 200 else None
        record("T09b", "Cashin : exécution", "POST /cashin/pay", "SUCCESSFULL", r, st == "SUCCESSFULL")
        r = get(f"/cashin/paymentstatus/{ci}", token)
        st = (r.json().get("data") or {}).get("status") if r.status_code == 200 else None
        record("T09c", "Cashin : statut", "GET /cashin/paymentstatus/{payToken}", "SUCCESSFULL", r, st == "SUCCESSFULL")

    # T10 — Cashout : réservé aux distributeurs (CICO), non concerné
    results.append({"N°": "T10", "Scénario": "Cashout complet (distributeurs)", "Endpoint": "/cashout/…",
                    "Résultat attendu": "PENDING en Sandbox", "HTTP": "", "Message": "", "Statut": "", "payToken": "",
                    "txnid": "", "orderId": "", "Résultat": "Non applicable",
                    "Commentaire": "WorldSoft Transfer n'est pas distributeur CICO", "Date et heure": now()})

    # T06 — opération non confirmée : on attend l'expiration du paiement T03 (jamais confirmé)
    if pt:
        print(f"T06 : attente de {FAIL_WAIT} s de l'expiration du paiement non confirmé…")
        deadline = time.time() + FAIL_WAIT
        r = None
        while time.time() < deadline:
            time.sleep(20)
            tk = get_token().json().get("access_token") or token
            r = get(f"/mp/paymentstatus/{pt}", tk)
            st = (r.json().get("data") or {}).get("status") if r.status_code == 200 else None
            if st and st != "PENDING":
                break
        st = (r.json().get("data") or {}).get("status") if r is not None and r.status_code == 200 else None
        record("T06", "Opération non confirmée", "GET /mp/paymentstatus/{payToken}", "FAILED, service non livré", r,
               st in ("FAILED", "EXPIRED", "CANCELLED"),
               "Service non livré (statut non SUCCESSFULL)" if st != "SUCCESSFULL" else "")

order = ["T01", "T02", "T03", "T04", "T05", "T06", "T07", "T08", "T09a", "T09b", "T09c", "T10"]
results.sort(key=lambda x: order.index(x["N°"]) if x["N°"] in order else 99)

out = HERE / "out"
out.mkdir(exist_ok=True)
stamp = datetime.now().strftime("%Y%m%d-%H%M")
cols = ["N°", "Scénario", "Endpoint", "Résultat attendu", "HTTP", "Statut", "Message", "payToken", "txnid", "orderId",
        "Résultat", "Commentaire", "Date et heure"]

md = [f"# Cahier de tests Sandbox Orange Money — WorldSoft Transfer (PWFINTECH)", "",
      f"Environnement : Sandbox · URL : {BASE_URL} · Date : {now()}", "",
      "| " + " | ".join(cols) + " |", "|" + "---|" * len(cols)]
for row in results:
    md.append("| " + " | ".join(str(row.get(c, "")).replace("|", "/") for c in cols) + " |")
(out / f"cahier-tests-orange-money-{stamp}.md").write_text("\n".join(md) + "\n")

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Tests Sandbox"
    ws.append(["Cahier de tests Sandbox Orange Money — WorldSoft Transfer (PWFINTECH)"])
    ws.append([f"Environnement : Sandbox · {BASE_URL} · {now()}"])
    ws.append([])
    ws.append(cols)
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="FF7900")
    for row in results:
        ws.append([row.get(c, "") for c in cols])
    for col, width in zip("ABCDEFGHIJKLM", [6, 34, 32, 30, 7, 13, 40, 32, 24, 22, 14, 40, 19]):
        ws.column_dimensions[col].width = width
    wb.save(out / f"cahier-tests-orange-money-{stamp}.xlsx")
except ImportError:
    pass

print(f"\nCahier écrit dans {out}")
